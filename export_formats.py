"""
Ankur newsletter — alternate export formats (Word, website HTML, poster,
blog post), built from the exact same `stories` list and `issue.json`
config the PDF generator (generate.py) uses. None of these reinvent the
PDF's layout -- each is its own, format-appropriate shape, reusing
generate.py's markup_tokens()/_lead_sentence() so ++highlight++/**bold**
markup and the no-markup auto-bold fallback stay consistent everywhere a
story's note text is rendered.

Four entry points, each (config, stories, output_path) -> output_path,
mirroring render_pdf()'s own signature in generate.py:
    render_docx(config, stories, output_path)
    render_website_html(config, stories, output_path)
    render_poster(config, story, output_path)        # one story, not a list
    write_blog_post(config, stories, output_path)     # AI-assisted, see below
"""

import base64
import html
import mimetypes
import textwrap
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont, ImageOps

from generate import _lead_sentence, markup_tokens
from styles import get_style, DEFAULT_STYLE_ID


# --- shared helpers --------------------------------------------------------

def _hex_rgb255(hex_color: str) -> tuple:
    hex_color = hex_color.lstrip("#")
    return tuple(int(hex_color[i:i + 2], 16) for i in (0, 2, 4))


def _float_rgb255(rgb: tuple) -> tuple:
    return tuple(round(c * 255) for c in rgb)


def _plain_text(text: str) -> str:
    """Strips ++highlight++/**bold** markup down to plain words, for
    contexts (poster pull-quotes) that can't render rich text."""
    return "".join(t for t, _ in markup_tokens(text))


# --- Word (.docx) export ----------------------------------------------------

def _add_markup_paragraph(doc: Document, para_text: str, highlight_hex: str):
    """Adds one Word paragraph from a note's paragraph text, honoring the
    same markup as generate.py's format_article(): explicit ++highlight++/
    **bold** markup if present, else auto-bold the lead sentence."""
    p = doc.add_paragraph()
    if "**" in para_text or "++" in para_text:
        for chunk, kind in markup_tokens(para_text):
            if not chunk:
                continue
            run = p.add_run(chunk)
            if kind in ("bold", "highlight"):
                run.bold = True
            if kind == "highlight":
                run.font.color.rgb = RGBColor(*_hex_rgb255(highlight_hex))
        return
    lead = _lead_sentence(para_text)
    rest = para_text[len(lead):]
    if lead:
        run = p.add_run(lead)
        run.bold = True
    if rest:
        p.add_run(rest)


def render_docx(config: dict, stories: list, output_path: Path) -> Path:
    """stories: list of (heading, note, photo_paths), same shape render_pdf() takes.
    A structured document (real headings/paragraphs/images), not a clone of
    the PDF's visual layout -- the point of a Word export is that someone
    can keep editing it."""
    style = get_style(config.get("style", DEFAULT_STYLE_ID))
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()

    title = doc.add_heading(f"{config.get('org_small', '')} \u2014 {config.get('title_big', 'NEWSLETTER')}".strip(" \u2014"), level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_range = (config.get("date_range") or "").strip()
    if date_range:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(date_range)
        run.italic = True

    for heading, note, photos in stories:
        if heading:
            doc.add_heading(heading, level=2)
        note = (note or "").strip()
        if note:
            for para in note.split("\n\n"):
                para = para.strip()
                if para:
                    _add_markup_paragraph(doc, para, style["highlight_hex"])
        for photo in photos:
            try:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run()
                run.add_picture(str(photo), width=Inches(6.0))
            except Exception:
                continue
        doc.add_paragraph()  # spacer between stories

    doc.save(str(output_path))
    return output_path


# --- website HTML export ----------------------------------------------------

def _markup_to_html(para_text: str, highlight_hex: str) -> str:
    out = []
    for chunk, kind in markup_tokens(para_text):
        escaped = html.escape(chunk).replace("\n", "<br>")
        if kind == "highlight":
            out.append(f'<b style="color:{highlight_hex}">{escaped}</b>')
        elif kind == "bold":
            out.append(f"<b>{escaped}</b>")
        else:
            out.append(escaped)
    return "".join(out)


def _story_html(heading: str, note: str, photos: list, highlight_hex: str) -> str:
    parts = [f"<article>"]
    if heading:
        parts.append(f"<h2>{html.escape(heading)}</h2>")
    note = (note or "").strip()
    if note:
        for para in note.split("\n\n"):
            para = para.strip()
            if not para:
                continue
            if "**" in para or "++" in para:
                body = _markup_to_html(para, highlight_hex)
            else:
                lead = _lead_sentence(para)
                rest = para[len(lead):]
                body = f"<b>{html.escape(lead)}</b>{html.escape(rest)}"
            parts.append(f"<p>{body}</p>")
    for photo in photos:
        try:
            data = Path(photo).read_bytes()
            mime = mimetypes.guess_type(str(photo))[0] or "image/jpeg"
            b64 = base64.b64encode(data).decode("ascii")
            parts.append(f'<img src="data:{mime};base64,{b64}" alt="{html.escape(heading)}">')
        except Exception:
            continue
    parts.append("</article>")
    return "\n".join(parts)


WEBSITE_TEMPLATE = """<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>{title}</title>
<style>
  body {{ font-family: Georgia, 'Times New Roman', serif; background:{cream}; color:{ink};
         max-width: 760px; margin: 0 auto; padding: 32px 20px 80px; line-height: 1.6; }}
  header {{ text-align:center; margin-bottom: 40px; }}
  header .org {{ letter-spacing: .12em; font-size: 13px; font-weight: 700; color:{accent}; text-transform: uppercase; }}
  header h1 {{ font-size: 36px; margin: 6px 0; }}
  header .date {{ font-size: 14px; color:{ink}; opacity: .7; }}
  article {{ margin: 0 0 48px; padding-bottom: 32px; border-bottom: 1px solid rgba(0,0,0,.1); }}
  article h2 {{ font-size: 22px; margin-bottom: 10px; }}
  article img {{ max-width: 100%; height: auto; border-radius: 8px; margin: 10px 0; display: block; }}
  article p {{ margin: 0 0 12px; }}
</style>
</head>
<body>
<header>
  <div class="org">{org_small}</div>
  <h1>{title_big}</h1>
  <div class="date">{date_range}</div>
</header>
<main>
{stories_html}
</main>
</body>
</html>
"""


def render_website_html(config: dict, stories: list, output_path: Path) -> Path:
    """A clean, self-contained (images inlined as base64) HTML page built
    from the same stories -- ready to paste into a website's CMS or drop
    onto static hosting as-is."""
    style = get_style(config.get("style", DEFAULT_STYLE_ID))
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    stories_html = "\n".join(
        _story_html(heading, note, photos, style["highlight_hex"]) for heading, note, photos in stories
    )
    page = WEBSITE_TEMPLATE.format(
        title=html.escape(f"{config.get('org_small', '')} {config.get('title_big', 'Newsletter')}".strip()),
        cream="#%02x%02x%02x" % _float_rgb255(style["cream"]),
        ink="#%02x%02x%02x" % _float_rgb255(style["ink"]),
        accent=style["accent_hex"],
        org_small=html.escape(config.get("org_small", "")),
        title_big=html.escape(config.get("title_big", "Newsletter")),
        date_range=html.escape(config.get("date_range", "")),
        stories_html=stories_html,
    )
    output_path.write_text(page, encoding="utf-8")
    return output_path


# --- poster (single story) --------------------------------------------------

POSTER_W, POSTER_H = 1240, 1754  # ~A4 at 150dpi, portrait
POSTER_MARGIN = 70

# Raw TTF filenames (not reportlab font names) for the poster's own PIL
# rendering -- a separate, smaller mapping from styles.py's reportlab font
# registration, since PIL needs a file path, not a registered font name.
_POSTER_FONTS = {
    "ankur-classic": ("arialbd.ttf", "arial.ttf"),
    "broadsheet-editorial": ("timesbd.ttf", "times.ttf"),
    "bold-magazine": ("GOTHICB.TTF", "GOTHIC.TTF"),
    "tabloid-pop": ("impact.ttf", "arial.ttf"),
    "minimalist": ("segoeuib.ttf", "segoeui.ttf"),
    "formal-bulletin": ("georgiab.ttf", "georgia.ttf"),
    "pastel-community": ("Candarab.ttf", "Candara.ttf"),
    "sports-two-tone": ("ARIALNB.TTF", "ARIALN.TTF"),
}
_FONT_DIR = Path(r"C:\Windows\Fonts")


def _poster_font(style_id: str, bold: bool, size: int) -> ImageFont.FreeTypeFont:
    bold_file, regular_file = _POSTER_FONTS.get(style_id, ("arialbd.ttf", "arial.ttf"))
    filename = bold_file if bold else regular_file
    try:
        return ImageFont.truetype(str(_FONT_DIR / filename), size)
    except Exception:
        try:
            return ImageFont.load_default(size=size)
        except Exception:
            return ImageFont.load_default()


def _biggest_photo(photos: list):
    best, best_area = None, -1
    for p in photos:
        try:
            with Image.open(p) as im:
                area = im.size[0] * im.size[1]
            if area > best_area:
                best, best_area = p, area
        except Exception:
            continue
    return best


def _wrap_to_width(draw, text, font, max_width):
    words = text.split()
    lines, cur = [], ""
    for w in words:
        trial = f"{cur} {w}".strip()
        if draw.textlength(trial, font=font) <= max_width or not cur:
            cur = trial
        else:
            lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines


def render_poster(config: dict, story: tuple, output_path: Path) -> Path:
    """story: one (heading, note, photo_paths) tuple. A single-page
    highlight poster for one story, using the current style's own palette
    (cream background, ink for readable text, accent/highlight for color),
    not a shrunk-down newsletter block -- the headline and photo dominate.
    """
    heading, note, photos = story
    style_id = config.get("style", DEFAULT_STYLE_ID)
    style = get_style(style_id)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    cream = _float_rgb255(style["cream"])
    ink = _float_rgb255(style["ink"])
    accent = _hex_rgb255(style["accent_hex"])
    highlight = _hex_rgb255(style["highlight_hex"])

    img = Image.new("RGB", (POSTER_W, POSTER_H), cream)
    draw = ImageDraw.Draw(img)
    content_w = POSTER_W - 2 * POSTER_MARGIN
    y = POSTER_MARGIN

    # eyebrow: org name on an accent pill
    org_small = (config.get("org_small") or "").strip()
    if org_small:
        eyebrow_font = _poster_font(style_id, bold=True, size=26)
        pill_w = draw.textlength(org_small, font=eyebrow_font) + 48
        draw.rounded_rectangle([POSTER_MARGIN, y, POSTER_MARGIN + pill_w, y + 52], radius=26, fill=accent)
        draw.text((POSTER_MARGIN + 24, y + 11), org_small, font=eyebrow_font, fill=(255, 255, 255))
        y += 52 + 30

    # headline
    heading_text = (heading or "This story").strip()
    headline_size = 76
    headline_font = _poster_font(style_id, bold=True, size=headline_size)
    lines = _wrap_to_width(draw, heading_text, headline_font, content_w)
    while len(lines) > 4 and headline_size > 40:
        headline_size -= 6
        headline_font = _poster_font(style_id, bold=True, size=headline_size)
        lines = _wrap_to_width(draw, heading_text, headline_font, content_w)
    for line in lines:
        draw.text((POSTER_MARGIN, y), line, font=headline_font, fill=ink)
        y += int(headline_size * 1.18)
    y += 26

    # accent rule under the headline
    draw.rectangle([POSTER_MARGIN, y, POSTER_MARGIN + 160, y + 6], fill=accent)
    y += 36

    # photo
    photo = _biggest_photo(photos)
    if photo:
        remaining_h = POSTER_H - POSTER_MARGIN - y - 260  # leave room for quote + footer
        if remaining_h > 200:
            try:
                with Image.open(photo) as im:
                    im = ImageOps.exif_transpose(im).convert("RGB")
                    fitted = ImageOps.fit(im, (content_w, remaining_h), Image.LANCZOS, centering=(0.5, 0.4))
                img.paste(fitted, (POSTER_MARGIN, y))
                y += remaining_h + 36
            except Exception:
                pass

    # pull-quote: first sentence(ish) of the note, markup stripped
    plain_note = _plain_text(note or "").strip()
    if plain_note:
        quote = _lead_sentence(plain_note)
        if len(quote) > 220:
            quote = quote[:217].rsplit(" ", 1)[0] + "..."
        quote_font = _poster_font(style_id, bold=False, size=34)
        quote_lines = _wrap_to_width(draw, f"\u201c{quote}\u201d", quote_font, content_w)[:4]
        for line in quote_lines:
            draw.text((POSTER_MARGIN, y), line, font=quote_font, fill=tuple(min(255, c + 0) for c in ink))
            y += 44

    # footer bar
    footer_h = 90
    draw.rectangle([0, POSTER_H - footer_h, POSTER_W, POSTER_H], fill=accent)
    footer_font = _poster_font(style_id, bold=True, size=28)
    footer_text = " \u00b7 ".join(t for t in (org_small, (config.get("date_range") or "").strip()) if t)
    if footer_text:
        draw.text((POSTER_MARGIN, POSTER_H - footer_h + 28), footer_text, font=footer_font, fill=(255, 255, 255))

    img.save(output_path, "PNG")
    return output_path


# --- AI blog post ------------------------------------------------------------

BLOG_SYSTEM_PROMPT = """You turn a school's internal newsletter story list into one cohesive blog \
post for the school's public website.

Rules:
- Use ONLY facts present in the stories given below. Never invent names, dates, numbers, or details \
not present in the source material.
- Write flowing narrative prose in paragraphs, not bullet points or a list of separate blurbs -- \
weave the stories together into one engaging piece a parent or visitor would enjoy reading, with a \
natural opening and closing, not just the stories concatenated back to back.
- This is a longer, warmer, more narrative register than a newsletter blurb -- take the space to set \
scenes and give context, while staying strictly grounded in what the stories actually say.
- Output valid Markdown: a single '# ' title for the whole post, then flowing paragraphs. You may use \
'## ' subheadings if it helps the post's structure, and **bold** sparingly for genuinely standout facts.
- Output ONLY the Markdown. No preamble, no explanation, no code fences.
"""


def _stories_as_material(stories: list) -> str:
    chunks = []
    for heading, note, _photos in stories:
        plain = _plain_text(note or "")
        chunks.append(f"Story: {heading or '(no heading)'}\n{plain}")
    return "\n\n".join(chunks)


def _fallback_blog_markdown(config: dict, stories: list) -> str:
    """Deterministic fallback used with no ANTHROPIC_API_KEY, or if the AI
    call fails -- a plain concatenation, not a narrative, but never blocks
    the export on a missing/broken API key (same discipline as summarize.py)."""
    title = config.get("title_big", "Newsletter")
    org = config.get("org_small", "")
    lines = [f"# {org} {title}".strip(), ""]
    date_range = (config.get("date_range") or "").strip()
    if date_range:
        lines += [f"*{date_range}*", ""]
    for heading, note, _photos in stories:
        if heading:
            lines.append(f"## {heading}")
        plain = _plain_text(note or "").strip()
        if plain:
            lines.append(plain)
        lines.append("")
    return "\n".join(lines).strip() + "\n"


def write_blog_post(config: dict, stories: list, output_path: Path) -> Path:
    """AI-assisted narrative blog post woven from all current stories.
    Requires ANTHROPIC_API_KEY (see summarize.py / .env.example) for real
    results; without one, or if the call fails, falls back to a plain
    deterministic concatenation rather than erroring."""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    markdown = None
    try:
        import summarize  # reuses the same lazy Anthropic client + .env loading
        client = summarize._get_client()
        if client is not None:
            response = client.messages.create(
                model="claude-opus-5-5",
                max_tokens=2000,
                output_config={"effort": "medium"},
                system=BLOG_SYSTEM_PROMPT,
                messages=[{"role": "user", "content": _stories_as_material(stories)}],
            )
            text = "".join(b.text for b in response.content if b.type == "text").strip()
            if text:
                markdown = text
    except Exception as e:
        print(f"WARNING: AI blog generation failed ({e}) -- using plain concatenation.")

    if not markdown:
        markdown = _fallback_blog_markdown(config, stories)

    output_path.write_text(markdown, encoding="utf-8")
    return output_path
