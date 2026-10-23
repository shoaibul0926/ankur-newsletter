"""
Ankur newsletter — Phase 2: deck extraction.

Pulls a heading, note text, and photos out of each slide (.pptx), page
(.pdf), or Heading-styled section (.docx) of an uploaded file, and writes
them into stories/NN-.../ folders in the exact format generate.py already
reads (heading.txt, note.txt, photos/).

This is a MECHANICAL extraction step only — it faithfully pulls whatever
text and photos are on the slide/page. It does NOT trim verbose slide text
down to "just the main point" (that's an editorial judgment call, not a
layout one) — review the extracted note.txt files afterwards and shorten
them before running generate.py, the same way Ankur's real newsletter only
keeps the headline fact per story rather than the full raw slide content.

Usage:
    python extract.py path/to/deck.pptx
    python extract.py path/to/deck.pdf
    python extract.py path/to/report.docx

Writes into stories/, numbered to continue after whatever is already there,
so running it twice (e.g. once per term's deck) doesn't clobber earlier
manually-written stories.
"""

import re
import sys
from pathlib import Path

from PIL import Image

BASE_DIR = Path(__file__).resolve().parent
MIN_IMAGE_SIDE = 150  # px; filters out small logos/icons/bullet graphics

_BULLET_RE = re.compile(r"^[\s]*[•◦▪‣●○\-\*]+\s*")
_PAGE_NUM_RE = re.compile(r"^(page\s*)?\d+(\s*(of|/)\s*\d+)?$", re.IGNORECASE)


def _clean_line(line: str) -> str:
    line = _BULLET_RE.sub("", line.strip())
    return re.sub(r"\s{2,}", " ", line).strip()


def _clean_text(text: str) -> str:
    lines = []
    for raw in text.splitlines():
        line = _clean_line(raw)
        if line and not _PAGE_NUM_RE.match(line):
            lines.append(line)
    return "\n".join(lines)


def _save_if_real_photo(data: bytes, dest: Path) -> bool:
    dest.write_bytes(data)
    try:
        with Image.open(dest) as im:
            w, h = im.size
        if w < MIN_IMAGE_SIDE and h < MIN_IMAGE_SIDE:
            dest.unlink(missing_ok=True)
            return False
    except Exception:
        dest.unlink(missing_ok=True)
        return False
    return True


def extract_pptx(path: Path, img_out_dir: Path) -> list[dict]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    prs = Presentation(str(path))
    results = []

    for slide_idx, slide in enumerate(prs.slides, start=1):
        heading = ""
        body_lines = []
        images = []
        title_shape = slide.shapes.title
        slide_img_dir = img_out_dir / f"slide_{slide_idx}"

        for shape in slide.shapes:
            if shape.has_text_frame:
                text = "\n".join(p.text for p in shape.text_frame.paragraphs if p.text.strip()).strip()
                if not text:
                    pass
                elif shape is title_shape and not heading:
                    lines = text.splitlines()
                    heading = lines[0].strip()
                    if len(lines) > 1:
                        body_lines.append("\n".join(lines[1:]).strip())
                else:
                    body_lines.append(text)

            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    image = shape.image
                    slide_img_dir.mkdir(parents=True, exist_ok=True)
                    dest = slide_img_dir / f"img_{len(images) + 1}.{image.ext}"
                    if _save_if_real_photo(image.blob, dest):
                        images.append(dest)
                except Exception:
                    pass

        if not heading and body_lines:
            first_lines = body_lines.pop(0).splitlines()
            heading = first_lines[0].strip()
            if len(first_lines) > 1:
                body_lines.insert(0, "\n".join(first_lines[1:]))

        heading = _clean_line(heading)
        note = _clean_text("\n\n".join(body_lines))
        if heading or note or images:
            results.append({"heading": heading, "note": note, "images": images})

    return results


def extract_pdf(path: Path, img_out_dir: Path) -> list[dict]:
    import pymupdf as fitz

    doc = fitz.open(str(path))
    results = []

    for page_idx in range(doc.page_count):
        page = doc[page_idx]
        lines = [_clean_line(l) for l in page.get_text("text").splitlines() if l.strip()]
        lines = [l for l in lines if l and not _PAGE_NUM_RE.match(l)]
        heading = lines[0] if lines else ""
        note = "\n".join(lines[1:]).strip()

        images = []
        page_img_dir = img_out_dir / f"page_{page_idx + 1}"
        for img in page.get_images(full=True):
            try:
                base = doc.extract_image(img[0])
                page_img_dir.mkdir(parents=True, exist_ok=True)
                dest = page_img_dir / f"img_{len(images) + 1}.{base['ext']}"
                if _save_if_real_photo(base["image"], dest):
                    images.append(dest)
            except Exception:
                pass

        if heading or note or images:
            results.append({"heading": heading, "note": note, "images": images})

    doc.close()
    return results


def extract_docx(path: Path, img_out_dir: Path) -> list[dict]:
    """One story per Heading-styled paragraph; plain paragraphs between
    headings become that story's note. Inline images are matched to the
    paragraph they actually appear in via the paragraph's own XML, so a
    photo under the wrong heading isn't a risk.
    """
    import docx

    A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
    R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"

    document = docx.Document(str(path))
    results: list[dict] = []
    heading, body_lines, images = "", [], []
    img_counter = 0

    def flush():
        nonlocal heading, body_lines, images
        heading_c = _clean_line(heading)
        note_c = _clean_text("\n".join(body_lines))
        if heading_c or note_c or images:
            results.append({"heading": heading_c, "note": note_c, "images": images})
        heading, body_lines, images = "", [], []

    for para in document.paragraphs:
        text = para.text.strip()
        is_heading = bool(para.style and para.style.name and para.style.name.lower().startswith("heading"))

        para_images = []
        for blip in para._p.findall(f".//{{{A_NS}}}blip"):
            rId = blip.get(f"{{{R_NS}}}embed")
            if not rId:
                continue
            try:
                part = para.part.related_parts[rId]
            except KeyError:
                continue
            img_counter += 1
            img_out_dir.mkdir(parents=True, exist_ok=True)
            dest = img_out_dir / f"img_{img_counter}{Path(part.partname).suffix or '.png'}"
            if _save_if_real_photo(part.blob, dest):
                para_images.append(dest)

        if is_heading and text:
            if heading or body_lines or images:
                flush()
            heading = text
            images.extend(para_images)
        else:
            if text:
                body_lines.append(text)
            images.extend(para_images)

    flush()
    return results


def extract_deck(path: Path, img_out_dir: Path) -> list[dict]:
    ext = path.suffix.lower()
    if ext == ".pptx":
        return extract_pptx(path, img_out_dir)
    if ext == ".pdf":
        return extract_pdf(path, img_out_dir)
    if ext == ".docx":
        return extract_docx(path, img_out_dir)
    raise ValueError(f"Unsupported file type: {ext} (expected .pptx, .pdf, or .docx)")


def _next_story_index(stories_dir: Path) -> int:
    existing = [p.name for p in stories_dir.iterdir() if p.is_dir()]
    nums = []
    for name in existing:
        m = re.match(r"^(\d+)-", name)
        if m:
            nums.append(int(m.group(1)))
    return (max(nums) + 1) if nums else 1


def write_stories(slides: list[dict], stories_dir: Path, deck_name: str) -> list[Path]:
    import shutil

    stories_dir.mkdir(parents=True, exist_ok=True)
    start = _next_story_index(stories_dir)
    written = []

    for offset, slide in enumerate(slides):
        idx = start + offset
        slug = re.sub(r"[^a-z0-9]+", "-", (slide["heading"] or deck_name).lower()).strip("-")[:40] or "story"
        folder = stories_dir / f"{idx:02d}-{slug}"
        folder.mkdir(parents=True, exist_ok=True)

        (folder / "heading.txt").write_text(slide["heading"] + "\n", encoding="utf-8")
        (folder / "note.txt").write_text(slide["note"] + "\n", encoding="utf-8")

        if slide["images"]:
            photos_dir = folder / "photos"
            photos_dir.mkdir(exist_ok=True)
            for i, img_path in enumerate(slide["images"], start=1):
                shutil.copy(img_path, photos_dir / f"img_{i}{img_path.suffix.lower()}")

        written.append(folder)

    return written


def main():
    if len(sys.argv) < 2:
        print("Usage: python extract.py path/to/deck.pptx|deck.pdf|report.docx")
        sys.exit(1)

    deck_path = Path(sys.argv[1]).resolve()
    if not deck_path.exists():
        print(f"File not found: {deck_path}")
        sys.exit(1)

    raw_out_dir = BASE_DIR / "web_uploads" / "extract" / deck_path.stem
    slides = extract_deck(deck_path, raw_out_dir)

    if not slides:
        print("No slides/pages with text or photos were found in that deck.")
        return

    written = write_stories(slides, BASE_DIR / "stories", deck_path.stem)

    print(f"Extracted {len(slides)} slide(s)/page(s) from {deck_path.name} into:")
    for folder in written:
        heading = (folder / "heading.txt").read_text(encoding="utf-8").strip()
        note = (folder / "note.txt").read_text(encoding="utf-8").strip()
        n_photos = len(list((folder / "photos").glob("*"))) if (folder / "photos").exists() else 0
        flag = "  <-- NEEDS A CAPTION (has photos, no note)" if n_photos and not note else ""
        print(f"  {folder.relative_to(BASE_DIR)}  —  \"{heading or '(no heading)'}\"  ({n_photos} photo(s)){flag}")
    print("\nReview/trim each note.txt down to just the main point before running generate.py —")
    print("extraction pulls the raw slide text as-is, it doesn't shorten it for you. Every photo")
    print("needs a caption explaining it — write one for any story flagged above.")


if __name__ == "__main__":
    main()
